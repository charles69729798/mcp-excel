
import os
import json
import tkinter as tk
from tkinter import filedialog
import google.generativeai as genai
from pypdf import PdfReader
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

# --- 1. 설정 ---
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "AIzaSyCNiN_y_8DAyuX9iFWgG6HK_mlEXlmsspQ")
PROMPT_FILE_PATH = "C:\\InsuranceProject\\insurance_analysis_prompt.txt"
OUTPUT_EXCEL_PATH = "C:\\InsuranceProject\\상품설계명세서.xlsx"
# JSON 결과를 저장할 경로 추가
OUTPUT_JSON_PATH = "C:\\InsuranceProject\\분석결과.json"

# --- 2. GUI 파일 선택 함수 ---
def get_pdf_path_from_gui():
    """GUI를 통해 사용자에게 PDF 파일 선택을 요청합니다."""
    root = tk.Tk()
    root.withdraw()  # 메인 Tk 창을 숨깁니다.
    
    print("파일 선택 대화상자를 엽니다. 분석할 PDF 파일을 선택해주세요...")
    
    file_path = filedialog.askopenfilename(
        title="분석할 PDF 사업방법서 파일을 선택하세요",
        initialdir=os.getcwd(), # 현재 작업 디렉토리에서 시작
        filetypes=[("PDF files", "*.pdf"), ("All files", "*.*")])
    
    if not file_path:
        print("파일이 선택되지 않았습니다. 작업을 중단합니다.")
        return None
    
    print(f"선택된 파일: {file_path}")
    return file_path

# --- 3. 파일 읽기 및 API 호출 함수 ---
def read_file_content(file_path):
    print(f"파일 읽는 중: {os.path.basename(file_path)}...")
    try:
        if file_path.lower().endswith(".pdf"):
            reader = PdfReader(file_path)
            text = ""
            for page in reader.pages:
                text += page.extract_text() or ""
            return text
        else:
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read()
    except Exception as e:
        print(f"파일 읽기 오류: {e}")
        return None

def analyze_with_gemini(api_key, prompt, document_text):
    print("Gemini API 호출하여 문서 분석 중...")
    if not api_key or api_key == "YOUR_API_KEY_HERE":
        print("오류: Gemini API 키가 설정되지 않았습니다.")
        return None
    try:
        genai.configure(api_key=api_key)
        model = genai.GenerativeModel('models/gemini-2.5-pro')
        
        full_prompt = f"{prompt}\n\n--- 분석 대상 문서 ---\n{document_text}"
        response = model.generate_content(full_prompt)
        
        json_start_index = response.text.find('{')
        json_end_index = response.text.rfind('}')
        if json_start_index == -1 or json_end_index == -1:
            print("오류: API 응답에서 유효한 JSON을 찾을 수 없습니다.")
            print(f"원본 응답: {response.text}")
            return None
        json_string = response.text[json_start_index : json_end_index + 1]
        return json.loads(json_string)
    except Exception as e:
        print(f"API 호출 중 오류: {e}")
        return None

# --- 4. 엑셀 생성 함수 ---
def create_spec_sheet(data, output_path):
    print(f"'{os.path.basename(output_path)}' 엑셀 파일 생성 중...")
    workbook = Workbook()
    ws = workbook.active
    ws.title = "입력"
    
    row = 1
    # Helper for writing sections
    def write_section(title, data_dict, start_row):
        ws[f'A{start_row}'] = title
        ws[f'A{start_row}'].font = Font(bold=True, size=14)
        current_row = start_row + 1
        for key, value in data_dict.items():
            ws[f'B{current_row}'] = key
            ws[f'C{current_row}'] = str(value)
            current_row += 1
        return current_row + 1

    # Product Info
    ws[f'A{row}'] = "상품 기본 정보"
    ws[f'A{row}'].font = Font(bold=True, size=16)
    row += 1
    ws[f'B{row}'] = "productName"
    ws[f'C{row}'] = data.get("productName", "")
    row += 1
    ws[f'B{row}'] = "productCategory"
    ws[f'C{row}'] = data.get("productCategory", "")
    row += 2

    # Plan Axes
    row = write_section("상품 판매 유형 (Plan Axes)", data.get("planAxes", {}), row)

    # Main Contract
    main_contract = data.get("mainContract", {})
    ws[f'A{row}'] = "주계약 (Main Contract)"
    ws[f'A{row}'].font = Font(bold=True, size=16)
    row += 1
    row = write_section("파라미터 (Params)", main_contract.get("params", {}), row)
    
    # Main Contract Rules
    ws[f'B{row}'] = "규칙 (Rules)"
    ws[f'B{row}'].font = Font(bold=True, size=12)
    row += 1
    for i, rule in enumerate(main_contract.get("rules", [])):
        ws[f'C{row}'] = f"Rule {i+1}"
        ws[f'D{row}'] = f'IF: {str(rule.get("if"))}'
        row += 1
        ws[f'D{row}'] = f'THEN: {str(rule.get("then"))}'
        row += 1
    row += 1

    # Riders
    ws[f'A{row}'] = "특약 (Riders)"
    ws[f'A{row}'].font = Font(bold=True, size=16)
    row += 1
    for rider in data.get("riders", []):
        rider_name = rider.get("riderName", "N/A")
        ws[f'B{row}'] = f"특약명: {rider_name}"
        ws[f'B{row}'].font = Font(bold=True, size=14)
        row += 1
        row = write_section("파라미터 (Params)", rider.get("params", {}), row)
        
        ws[f'C{row}'] = "규칙 (Rules)"
        ws[f'C{row}'].font = Font(bold=True, size=12)
        row += 1
        for i, rule in enumerate(rider.get("rules", [])):
            ws[f'D{row}'] = f"Rule {i+1}: {str(rule)}"
            row += 1
        row += 1

    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 40
    ws.column_dimensions['D'].width = 60

    workbook.save(output_path)
    print(f"성공! '{output_path}'에 상품 설계 명세서가 저장되었습니다.")

# --- 5. 메인 실행 블록 ---
if __name__ == "__main__":
    if not GEMINI_API_KEY or GEMINI_API_KEY == "YOUR_API_KEY_HERE":
        print("오류: 스크립트 상단의 GEMINI_API_KEY를 본인의 키로 수정해주세요.")
    else:
        # 텍스트 기반 테스트를 위해 GUI를 비활성화하고 하드코딩된 경로를 사용합니다.
        pdf_path = "C:\\InsuranceProject\\034_사업방법서_신한주니어큐브종합건강상해보험(무배당, 해약환급금 미지급형)_20240401.pdf"
        print(f"텍스트 기반 테스트 모드: {os.path.basename(pdf_path)} 파일을 사용합니다.")

        if pdf_path:
            prompt_content = read_file_content(PROMPT_FILE_PATH)
            doc_content = read_file_content(pdf_path)

            if prompt_content and doc_content:
                analysis_result = analyze_with_gemini(GEMINI_API_KEY, prompt_content, doc_content)
                
                if analysis_result:
                    # 엑셀 파일 생성
                    create_spec_sheet(analysis_result, OUTPUT_EXCEL_PATH)
                    
                    # JSON 파일 저장
                    print(f"분석 결과를 JSON 파일로 저장하는 중: {os.path.basename(OUTPUT_JSON_PATH)}...")
                    try:
                        with open(OUTPUT_JSON_PATH, 'w', encoding='utf-8') as f:
                            json.dump(analysis_result, f, ensure_ascii=False, indent=4)
                        print("JSON 파일 저장 완료.")
                    except Exception as e:
                        print(f"JSON 파일 저장 중 오류 발생: {e}")
