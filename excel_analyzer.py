
import openpyxl
import re
import os

EXCEL_FILE_PATH = "C:\\InsuranceProject\\main.xlsm"
MACRO_ANALYSIS_FILE = "C:\\InsuranceProject\\macro_analysis.txt"
REPORT_OUTPUT_FILE = "C:\\InsuranceProject\\excel_analysis_report.txt"

def read_sheet_data(file_path, sheet_name):
    """
    Reads values and formulas from a specified sheet in an Excel file.
    Returns a list of lists representing the sheet data.
    """
    try:
        workbook = openpyxl.load_workbook(file_path, data_only=False) # data_only=False to get formulas
        if sheet_name not in workbook.sheetnames:
            return None, f"Sheet '{sheet_name}' not found."
        
        sheet = workbook[sheet_name]
        data = []
        for row in sheet.iter_rows():
            row_data = []
            for cell in row:
                if cell.data_type == 'f': # 'f' for formula
                    row_data.append(f"={cell.value}")
                else:
                    row_data.append(cell.value)
            data.append(row_data)
        return data, None
    except Exception as e:
        return None, f"Error reading sheet '{sheet_name}': {e}"

def parse_vba_code(macro_analysis_file):
    """
    Parses the macro_analysis.txt file to extract individual VBA modules and their code.
    Returns a dictionary where keys are module names and values are their VBA code.
    """
    vba_modules = {}
    current_module = None
    current_code = []

    try:
        with open(macro_analysis_file, 'r', encoding='utf-8') as f:
            for line in f:
                module_header_match = re.match(r'--- Module: (.*), Type: (\d+) ---', line)
                if module_header_match:
                    if current_module and current_code:
                        vba_modules[current_module] = "".join(current_code).strip()
                    current_module = module_header_match.group(1)
                    current_code = []
                elif line.startswith('--- Button Macro Analysis ---') or line.startswith('--- VBA Code Extraction ---'):
                    if current_module and current_code:
                        vba_modules[current_module] = "".join(current_code).strip()
                    current_module = None # Reset for non-module sections
                    current_code = []
                elif current_module:
                    current_code.append(line)
            if current_module and current_code:
                vba_modules[current_module] = "".join(current_code).strip()
    except FileNotFoundError:
        return None, f"Macro analysis file not found: {macro_analysis_file}"
    except Exception as e:
        return None, f"Error parsing VBA code: {e}"
    
    return vba_modules, None

def analyze_vba_sheet_interactions(vba_modules, target_sheet_names):
    """
    Analyzes VBA code for interactions (read/write) with specified sheets.
    Returns a dictionary summarizing interactions for each target sheet.
    """
    interactions = {name: {"reads": [], "writes": [], "other": []} for name in target_sheet_names}

    for module_name, code in vba_modules.items():
        for sheet_name in target_sheet_names:
            # Pattern for reading from sheet
            read_pattern = re.compile(r'Worksheets\(\s*["\']' + re.escape(sheet_name) + r'["\']\s*\)\.(Range|Cells)\([^)]+\)', re.IGNORECASE)
            # Pattern for writing to sheet (assignment)
            write_pattern = re.compile(r'Worksheets\(\s*["\']' + re.escape(sheet_name) + r'["\']\s*\)\.(Range|Cells)\([^)]+\)\s*=', re.IGNORECASE)
            # Pattern for selecting sheet
            select_pattern = re.compile(r'Worksheets\(\s*["\']' + re.escape(sheet_name) + r'["\']\s*\)\.Select', re.IGNORECASE)
            # Pattern for activating sheet
            activate_pattern = re.compile(r'Worksheets\(\s*["\']' + re.escape(sheet_name) + r'["\']\s*\)\.Activate', re.IGNORECASE)

            for line in code.splitlines():
                if write_pattern.search(line):
                    interactions[sheet_name]["writes"].append(f"Module '{module_name}': {line.strip()}")
                elif read_pattern.search(line):
                    interactions[sheet_name]["reads"].append(f"Module '{module_name}': {line.strip()}")
                elif select_pattern.search(line) or activate_pattern.search(line):
                    interactions[sheet_name]["other"].append(f"Module '{module_name}': {line.strip()}")
    return interactions

def generate_report(s_calc_data, vba_modules, interactions):
    """Generates a structured report of the analysis."""
    report = []
    report.append("---" + " Excel Workbook Analysis Report ---")
    report.append(f"Analyzed File: {EXCEL_FILE_PATH}")
    report.append(f"VBA Code Source: {MACRO_ANALYSIS_FILE}")
    report.append("-" * 40)

    # S산출 Sheet Analysis
    report.append("\n---" + " 'S산출' Sheet Analysis ---")
    if s_calc_data:
        report.append("Purpose: Primary interface for insurance product calculations and displaying results.")
        report.append("Key Areas:")
        report.append("  - Input Area (A1:L4): Applied Interest Rate, Base Amount, References to other sheets.")
        report.append("  - Output/Result Area (B9:Q18): Net Premiums, Alpha values, Error checks.")
        report.append("  - Actuarial Table (R2:AD~): Life table calculations (lx, qx, Cx, Mx, Dx, Nx) and Mortality Rates.")
        report.append("\nSample Data/Formulas from 'S산출' (first 10 rows, first 10 columns):")
        for r_idx, row in enumerate(s_calc_data):
            if r_idx >= 10: break
            report.append(f"  Row {r_idx+1}: {row[:10]}")
        report.append("\nInteractions with VBA Macros:")
        if interactions.get("S산출"):
            if interactions["S산출"]["reads"]:
                report.append("  Reads from 'S산출':")
                for line in interactions["S산출"]["reads"]:
                    report.append(f"    - {line}")
            if interactions["S산출"]["writes"]:
                report.append("  Writes to 'S산출':")
                for line in interactions["S산출"]["writes"]:
                    report.append(f"    - {line}")
            if interactions["S산출"]["other"]:
                report.append("  Other interactions with 'S산출':")
                for line in interactions["S산출"]["other"]:
                    report.append(f"    - {line}")
        else:
            report.append("  No direct VBA interactions identified for 'S산출'.")
    else:
        report.append("Could not read 'S산출' sheet data.")

    # 규정체크 Sheet Analysis (Inferred)
    report.append("\n---" + " '규정체크' Sheet Analysis (Inferred from VBA) ---")
    report.append("Purpose: Serves as a central repository for defining and checking various insurance product rules and parameters.")
    report.append("  (Note: Actual sheet content could not be read directly due to tool limitations.)")
    report.append("\nInteractions with VBA Macros:")
    if interactions.get("규정체크"):
        if interactions["규정체크"]["reads"]:
            report.append("  Reads from '규정체크':")
            for line in interactions["규정체크"]["reads"]:
                report.append(f"    - {line}")
        if interactions["규정체크"]["writes"]:
            report.append("  Writes to '규정체크':")
            for line in interactions["규정체크"]["writes"]:
                report.append(f"    - {line}")
        if interactions["규정체크"]["other"]:
            report.append("  Other interactions with '규정체크':")
            for line in interactions["규정체크"]["other"]:
                report.append(f"    - {line}")
    else:
        report.append("  No direct VBA interactions identified for '규정체크'.")

    # INPUT Sheet Analysis (Inferred)
    report.append("\n---" + " 'INPUT' Sheet Analysis (Inferred from VBA) ---")
    report.append("Purpose: Acts as an intermediary or a temporary input buffer for macros to feed parameters into core calculation logic.")
    report.append("  (Note: Actual sheet content could not be read directly due to tool limitations.)")
    report.append("\nInteractions with VBA Macros:")
    if interactions.get("INPUT"):
        if interactions["INPUT"]["reads"]:
            report.append("  Reads from 'INPUT':")
            for line in interactions["INPUT"]["reads"]:
                report.append(f"    - {line}")
        if interactions["INPUT"]["writes"]:
            report.append("  Writes to 'INPUT':")
            for line in interactions["INPUT"]["writes"]:
                report.append(f"    - {line}")
        if interactions["INPUT"]["other"]:
            report.append("  Other interactions with 'INPUT':")
            for line in interactions["INPUT"]["other"]:
                report.append(f"    - {line}")
    else:
        report.append("  No direct VBA interactions identified for 'INPUT'.")

    report.append("\n---" + " Summary of Dynamic Relationships ---")
    report.append("1.  'S산출' serves as the central calculation and display sheet, containing actuarial tables and receiving outputs from various macros.")
    report.append("2.  '규정체크' (inferred) is a rule definition and parameter generation sheet. The '규정체크()' macro systematically iterates through product configurations defined here, potentially writing results back.")
    report.append("3.  'INPUT' (inferred) acts as a crucial staging area. Macros write specific parameter combinations to 'INPUT', which are then likely consumed by formulas on 'S산출' or other calculation logic.")
    report.append("4.  The '보험료및준비금_전산반영_모두실행()' macro (in Sheet20) orchestrates a large-scale data generation process, iterating through parameters, using '설정()', '제한()', '입력설정()' (which interacts with 'INPUT'), and writing comprehensive results to external text files.")
    report.append("\nTo install macros into new Excel files, one would need to:")
    report.append("  - Replicate the structure and formulas of 'S산출' (especially the actuarial tables).")
    report.append("  - Understand and adapt the logic from 'Module2.규정체크()' for rule definition and parameter iteration.")
    report.append("  - Understand and adapt the calculation logic from 'Sheet18.S1Check()' and 'Sheet18.alpha_Check()'.")
    report.append("  - Recreate the 'INPUT' sheet's role as a parameter buffer.")
    report.append("  - Adapt the comprehensive data generation logic from 'Sheet20.보험료및준비금_전산반영_모두실행()' if mass output is required.")
    
    return "\n".join(report)

def main():
    s_calc_data, s_calc_error = read_sheet_data(EXCEL_FILE_PATH, "S산출")
    if s_calc_error:
        print(f"Error: {s_calc_error}")
        s_calc_data = None # Ensure it's None if there was an error

    vba_modules, vba_error = parse_vba_code(MACRO_ANALYSIS_FILE)
    if vba_error:
        print(f"Error: {vba_error}")
        return

    target_sheets = ["S산출", "규정체크", "INPUT"]
    interactions = analyze_vba_sheet_interactions(vba_modules, target_sheets)

    report_content = generate_report(s_calc_data, vba_modules, interactions)
    
    try:
        with open(REPORT_OUTPUT_FILE, "w", encoding="utf-8") as f:
            f.write(report_content)
        print(f"Analysis report generated successfully: {REPORT_OUTPUT_FILE}")
    except Exception as e:
        print(f"Error writing report to file: {e}")

if __name__ == "__main__":
    main()
