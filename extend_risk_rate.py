
import openpyxl

FILE_PATH = "C:\\InsuranceProject\\신한주니어큐브종합건강상해보험(무배당, 해약환급금 미지급형)_주보험.xlsx"

def extend_risk_rate():
    """'위험률' 시트를 복사하고 확장하여 '위험률_확장' 시트를 생성합니다."""
    try:
        workbook = openpyxl.load_workbook(FILE_PATH)

        # 1. '위험률' 시트 읽기
        if "위험률" not in workbook.sheetnames:
            raise ValueError("'위험률' 시트를 찾을 수 없습니다.")
        risk_sheet = workbook["위험률"]

        # 2. '위험률_확장' 시트 생성 또는 초기화
        if "위험률_확장" in workbook.sheetnames:
            del workbook["위험률_확장"]
        extended_sheet = workbook.create_sheet("위험률_확장")

        # 3. 데이터 복사
        for row in risk_sheet.iter_rows():
            for cell in row:
                extended_sheet[cell.coordinate].value = cell.value

        # 4. 데이터 확장 (51세부터 120세까지)
        last_age_row = 53 # 50세 데이터가 있는 행 (3행부터 시작)
        last_male_rate = extended_sheet[f'B{last_age_row}'].value
        last_female_rate = extended_sheet[f'C{last_age_row}'].value

        for age in range(51, 121):
            row = last_age_row + (age - 50)
            extended_sheet.cell(row=row, column=1, value=age)
            extended_sheet.cell(row=row, column=2, value=last_male_rate)
            extended_sheet.cell(row=row, column=3, value=last_female_rate)

        # 5. 파일 저장
        workbook.save(FILE_PATH)
        print("'위험률_확장' 시트가 성공적으로 생성되었습니다.")

    except Exception as e:
        print(f"오류 발생: {e}")

if __name__ == "__main__":
    extend_risk_rate()
