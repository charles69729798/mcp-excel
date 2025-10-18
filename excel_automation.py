
import os
import json
import google.generativeai as genai
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from pypdf import PdfReader

# --- 1. 설정 ---
# 본인의 Gemini API 키를 아래에 입력하거나, 환경 변수에 'GEMINI_API_KEY'로 설정하세요.
# 예: "YOUR_API_KEY_HERE"
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "AIzaSyCNiN_y_8DAyuX9iFWgG6HK_mlEXlmsspQ")

# 분석할 PDF 파일과 생성할 엑셀 파일의 경로
PDF_FILE_PATH = "C:\\InsuranceProject\\034_사업방법서_신한주니어큐브종합건강상해보험(무배당, 해약환급금 미지급형)_20240401.pdf"
EXCEL_FILE_PATH = "C:\\InsuranceProject\\신한주니어큐브_보험료산출_자동생성.xlsx"


def extract_text_from_pdf(pdf_path):
    """PDF 파일에서 텍스트를 추출합니다."""
    print(f"'{pdf_path}' 파일에서 텍스트를 추출하는 중...")
    try:
        reader = PdfReader(pdf_path)
        text = ""
        for page in reader.pages:
            text += page.extract_text() or ""
        print("텍스트 추출 완료.")
        return text
    except FileNotFoundError:
        print(f"오류: PDF 파일을 찾을 수 없습니다. 경로를 확인하세요: {pdf_path}")
        return None
    except Exception as e:
        print(f"PDF 처리 중 오류 발생: {e}")
        return None

def get_analysis_prompt(document_text):
    """Gemini API에 보낼 프롬프트를 생성합니다."""
    return f'''
당신은 보험 상품 분석 전문가입니다. 주어진 '사업방법서' 텍스트를 분석하여, 사용자가 엑셀에서 보험료 계산을 시뮬레이션하는 데 필요한 입력 항목들을 추출하고, 이를 JSON 형식으로 구조화해주세요.

**분석 규칙:**
1.  **'가입 유형'**: '어린이형'과 '태아형'을 값으로 추출합니다.
2.  **'상품 종류'**: '해약환급금 미지급형'과 '일반형'을 값으로 추출합니다.
3.  **'담보명'**: '주계약'과 모든 '종속 특약'의 명칭을 추출합니다.
4.  **'보험 기간' 및 '납입 기간'**: 각 담보, 가입유형, 상품종류에 따른 '보험기간'과 '보험료 납입기간'의 선택 가능한 값들을 모두 추출합니다.
5.  **'갱신 여부'**: 담보명에 '갱신형' 단어가 있으면 '갱신형', 없으면 '비갱신형'으로 분류합니다.
6.  **'출력값 매핑'**: '입력 방식'이 '드롭다운'인 항목에 대해, 각 '선택 가능 값'이 어떤 숫자/참조 값에 해당하는지 매핑 정보를 생성합니다. (예: '10년납'은 10, '30세만기'는 30). 이 매핑 정보가 없으면 null로 설정합니다.

**출력 형식 (JSON):**
- 반드시 다음 키를 포함하는 JSON 객체의 배열(list of objects)로 응답해주세요.
- 각 객체는 엑셀의 한 행(row)에 해당합니다.
- `["구분", "항목", "입력 방식", "선택 가능 값", "출력값 매핑", "관련 근거"]`
- "입력 방식"은 "드롭다운", "숫자 입력", "자동 표시" 중 하나여야 합니다.
- "선택 가능 값"은 "드롭다운" 방식일 경우 쉼표(,)로 구분된 문자열이어야 합니다.
- "출력값 매핑"은 JSON 객체(object) 형식이어야 합니다.

**예시 JSON 객체:**
{{
    "구분": "주계약",
    "항목": "납입 기간",
    "입력 방식": "드롭다운",
    "선택 가능 값": "10년납,20년납,30년납",
    "출력값 매핑": {{
        "10년납": 10,
        "20년납": 20,
        "30년납": 30
    }},
    "관련 근거": "사업방법서 2.가"
}}

이제 아래의 사업방법서 텍스트를 분석하여 완전한 JSON 배열을 생성해주세요.

--- 문서 텍스트 시작 ---
{document_text}
--- 문서 텍스트 끝 ---
'''

def analyze_document_with_gemini(api_key, document_text):
    """Gemini API를 호출하여 문서 분석을 요청하고, 구조화된 데이터를 받습니다."""
    if not api_key or api_key == "YOUR_API_KEY_HERE":
        print("오류: Gemini API 키가 설정되지 않았습니다. 스크립트 상단의 GEMINI_API_KEY를 수정하세요.")
        return None

    print("Gemini API 호출하여 문서 분석 중... (시간이 걸릴 수 있습니다)")
    try:
        genai.configure(api_key=api_key)
        model = genai.GenerativeModel('models/gemini-2.5-pro')
        
        prompt = get_analysis_prompt(document_text)
        
        response = model.generate_content(prompt)
        
        # API 응답에서 JSON 부분만 추출
        # 모델이 응답 앞뒤에 대화형 텍스트를 추가하는 경우가 있으므로, JSON 배열의 시작 '['과 끝 ']'을 찾습니다.
        json_start_index = response.text.find('[')
        json_end_index = response.text.rfind(']') # rfind()는 뒤에서부터 문자를 찾습니다.
        
        if json_start_index == -1 or json_end_index == -1:
            print("오류: API 응답에서 유효한 JSON 배열을 찾을 수 없습니다.")
            return None
        
        # JSON 시작과 끝 인덱스를 사용하여 정확한 JSON 문자열 부분만 추출합니다.
        json_string = response.text[json_start_index : json_end_index + 1]
        cleaned_response = json_string.strip()

        print("--- JSON 데이터 출력 ---")
        print(cleaned_response)
        print("-----------------------")

        print("API 분석 완료.")
        return json.loads(cleaned_response)
    except Exception as e:
        print(f"Gemini API 호출 중 오류 발생: {e}")
        print("API 키가 유효한지, 인터넷 연결이 정상인지 확인하세요.")
        return None

def create_excel_with_dropdowns(data, excel_path):
    """분석된 데이터로 드롭다운과 VLOOKUP 수식이 포함된 엑셀 파일을 생성합니다."""
    if not data:
        print("엑셀 파일을 생성할 데이터가 없습니다.")
        return

    print(f"'{excel_path}' 엑셀 파일 생성 중...")
    
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "입력"

    # --- 2단계: 데이터 매핑용 시트 생성 ---
    mapping_sheet = workbook.create_sheet(title="data_mapping")
    mapping_sheet.sheet_state = 'hidden' # 시트 숨기기
    mapping_row_counter = 1 # 매핑 시트의 행 카운터

    # 새로운 구조의 헤더 추가
    headers = ["구분", "항목", "입력값", "비고 / 선택 옵션", "관련 근거", "참조 값"]
    sheet.append(headers)

    # 데이터 추가 및 새로운 구조에 맞게 드롭다운/수식 설정
    for row_data in data:
        notes = ""
        if row_data.get("입력 방식") == "드롭다운":
            notes = f'선택: {row_data.get("선택 가능 값", "")}'
        elif row_data.get("입력 방식") == "숫자 입력":
            notes = "사용자 직접 입력"

        # --- 3단계: VLOOKUP 수식 추가 ---
        formula = ""
        mapping_info = row_data.get("출력값 매핑")

        if mapping_info:
            # 매핑 데이터를 mapping_sheet에 기록
            start_map_row = mapping_row_counter
            for key, value in mapping_info.items():
                mapping_sheet.cell(row=mapping_row_counter, column=1, value=key)
                mapping_sheet.cell(row=mapping_row_counter, column=2, value=value)
                mapping_row_counter += 1
            end_map_row = mapping_row_counter - 1
            
            # VLOOKUP 수식 생성
            # IFERROR를 사용하여 C열이 비어있을 때 F열도 비워둠
            current_row_index = sheet.max_row + 1
            lookup_cell = f"C{current_row_index}"
            formula = f'=IFERROR(VLOOKUP({lookup_cell}, data_mapping!A{start_map_row}:B{end_map_row}, 2, FALSE), "")'

        row_to_append = [
            row_data.get("구분"),
            row_data.get("항목"),
            "",  # C열: 입력값을 위한 빈 셀
            notes,
            row_data.get("관련 근거"),
            formula # F열: 수식
        ]
        sheet.append(row_to_append)
        
        row_idx = sheet.max_row
        
        if row_data.get("입력 방식") == "드롭다운":
            dropdown_values = row_data.get("선택 가능 값", "")
            if dropdown_values:
                target_cell = f"C{row_idx}"
                dv = DataValidation(type="list", formula1=f'"{dropdown_values}"', allow_blank=True)
                sheet.add_data_validation(dv)
                dv.add(target_cell)

    # 컬럼 너비 자동 조절
    for i, column_cells in enumerate(sheet.columns):
        if i != 2: # C열 제외
            max_length = 0
            column = column_cells[0].column_letter
            for cell in column_cells:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2) * 1.2
            sheet.column_dimensions[column].width = adjusted_width

    workbook.save(excel_path)
    print(f"성공! '{excel_path}'에 드롭다운과 수식이 포함된 엑셀 파일이 저장되었습니다.")


def main():
    """메인 실행 함수"""
    # 1. PDF에서 텍스트 추출
    document_text = extract_text_from_pdf(PDF_FILE_PATH)
    if not document_text:
        return

    # 2. Gemini로 문서 분석
    structured_data = analyze_document_with_gemini(GEMINI_API_KEY, document_text)
    if not structured_data:
        return
        
    # 3. 엑셀 파일 생성
    create_excel_with_dropdowns(structured_data, EXCEL_FILE_PATH)


if __name__ == "__main__":
    main()
