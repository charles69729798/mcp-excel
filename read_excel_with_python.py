import openpyxl

FILE_PATH = "C:\\InsuranceProject\\신한주니어큐브종합건강상해보험(무배당, 해약환급금 미지급형)_주보험_vba.xlsm"
SHEET_NAME = "s산출"
OUTPUT_FILE = "C:\\InsuranceProject\\s_sanchul_sheet_content.txt"

def read_sheet_data():
    """openpyxl을 사용하여 엑셀 시트의 모든 데이터를 읽어 파일에 씁니다."""
    try:
        workbook = openpyxl.load_workbook(FILE_PATH, data_only=True)
        
        if SHEET_NAME not in workbook.sheetnames:
            with open(OUTPUT_FILE, 'w', encoding='utf-8') as f:
                f.write(f"오류: '{SHEET_NAME}' 시트를 찾을 수 없습니다.\n")
            return

        sheet = workbook[SHEET_NAME]
        
        with open(OUTPUT_FILE, 'w', encoding='utf-8') as f:
            f.write(f"--- '{SHEET_NAME}' 시트 데이터 ---\n")
            has_data = False
            for row in sheet.iter_rows():
                row_values = [str(cell.value) if cell.value is not None else '' for cell in row]
                if any(val for val in row_values):
                    f.write(str(row_values) + '\n')
                    has_data = True
            
            if not has_data:
                f.write("시트에 데이터가 없습니다.\n")
            
            f.write("---------------------------\n")

    except FileNotFoundError:
        with open(OUTPUT_FILE, 'w', encoding='utf-8') as f:
            f.write(f"오류: 파일을 찾을 수 없습니다 - {FILE_PATH}\n")
    except Exception as e:
        with open(OUTPUT_FILE, 'w', encoding='utf-8') as f:
            f.write(f"엑셀 파일 처리 중 오류 발생: {e}\n")

if __name__ == "__main__":
    read_sheet_data()
